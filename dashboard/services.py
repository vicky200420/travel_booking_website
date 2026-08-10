"""Dashboard — business logic.

Single place for the derived statistics, membership/loyalty calculations,
wishlist operations and preferences so views stay thin.
"""
from collections import OrderedDict
from datetime import datetime, time, timedelta
from decimal import Decimal
from html import escape

from django.contrib.auth import get_user_model
from django.contrib.auth.password_validation import validate_password
from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import ValidationError
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.utils import timezone

from bookings.models import Booking
from flights.models import Flight
from hotels.models import Hotel
from notifications.models import Notification
from packages.models import TravelPackage
from payments.models import Payment
from reviews.models import Review

from .models import UserProfile, WishlistItem

User = get_user_model()

POINTS_PER_100 = Decimal('100')

MEMBERSHIP_TIERS = [
    (Decimal('500000'), 'Platinum', 'bi-gem', '#06B6D4'),
    (Decimal('200000'), 'Gold', 'bi-award', '#F59E0B'),
    (Decimal('50000'), 'Silver', 'bi-star-fill', '#94A3B8'),
    (Decimal('0'), 'Bronze', 'bi-shield-fill-check', '#B45309'),
]

WISHLIST_MODELS = {
    'hotel': ('hotels', 'hotel'),
    'flight': ('flights', 'flight'),
    'package': ('packages', 'travelpackage'),
}


class DashboardService:

    # ------------------------------------------------------------ profile

    @staticmethod
    def get_profile(user):
        profile, _ = UserProfile.objects.get_or_create(user=user)
        return profile.ensure_defaults()

    # -------------------------------------------------------- membership

    @staticmethod
    def total_spent(user):
        """Sum of successfully paid transactions for ``user``."""
        return (
            Payment.objects.filter(user=user, status=Payment.Status.SUCCESS)
            .aggregate(total=Sum('amount'))['total'] or Decimal('0')
        )

    @staticmethod
    def loyalty_points(user):
        """1 reward point per 100 of confirmed spend (placeholder formula)."""
        return int(DashboardService.total_spent(user) / POINTS_PER_100)

    @staticmethod
    def membership_level(user):
        """Tier derived from lifetime confirmed spend."""
        spent = DashboardService.total_spent(user)
        for threshold, label, icon, color in MEMBERSHIP_TIERS:
            if spent >= threshold:
                return {'key': label.lower(), 'label': label, 'icon': icon, 'color': color}
        return MEMBERSHIP_TIERS[-1]

    # ------------------------------------------------------------ stats

    @staticmethod
    def stats(user):
        now = timezone.now().date()
        bookings = Booking.objects.filter(user=user)

        upcoming = bookings.filter(
            status__in=[Booking.Status.PENDING, Booking.Status.CONFIRMED],
            travel_start_date__gte=now,
        ).count()

        return {
            'total_trips': bookings.count(),
            'upcoming_trips': upcoming,
            'completed_trips': bookings.filter(status=Booking.Status.COMPLETED).count(),
            'cancelled_trips': bookings.filter(status=Booking.Status.CANCELLED).count(),
            'total_spent': DashboardService.total_spent(user),
            'reward_points': DashboardService.loyalty_points(user),
            'unread_notifications': Notification.objects.filter(
                user=user, is_read=False
            ).count(),
        }

    # ----------------------------------------------------------- bookings

    @staticmethod
    def upcoming_bookings(user, limit=4):
        now = timezone.now().date()
        return (
            Booking.objects.filter(
                user=user,
                status__in=[Booking.Status.PENDING, Booking.Status.CONFIRMED],
                travel_start_date__gte=now,
            )
            .select_related('flight', 'hotel', 'package')
            .order_by('travel_start_date')[:limit]
        )

    @staticmethod
    def booking_queryset(user):
        return (
            Booking.objects.filter(user=user)
            .select_related('flight', 'hotel', 'package')
        )

    # ---------------------------------------------------------- payments

    @staticmethod
    def payment_queryset(user):
        return (
            Payment.objects.filter(user=user)
            .select_related('booking')
        )

    @staticmethod
    def monthly_spend(user, months=6):
        """Confirmed spend aggregated per calendar month (oldest → newest)."""
        since = timezone.now() - timezone.timedelta(days=months * 31)
        payments = Payment.objects.filter(
            user=user, status=Payment.Status.SUCCESS, paid_at__gte=since
        )
        buckets = OrderedDict()
        for i in range(months - 1, -1, -1):
            d = (timezone.now() - timezone.timedelta(days=30 * i)).date()
            buckets[(d.year, d.month)] = Decimal('0')

        for p in payments:
            if p.paid_at is None:
                continue
            key = (p.paid_at.year, p.paid_at.month)
            if key in buckets:
                buckets[key] += p.amount

        labels = [f'{d.month:02d}/{d.year}' for d in _bucket_dates(buckets)]
        values = [float(buckets[k]) for k in buckets]
        return labels, values

    # ----------------------------------------------------------- reviews

    @staticmethod
    def review_queryset(user):
        return (
            Review.objects.filter(user=user)
            .select_related('hotel', 'package', 'booking')
        )

    # ------------------------------------------------------ notifications

    @staticmethod
    def notification_queryset(user):
        return Notification.objects.filter(user=user)

    # ----------------------------------------------------------- wishlist

    @staticmethod
    def wishlist_queryset(user):
        return (
            WishlistItem.objects.filter(user=user)
            .select_related('content_type')
        )

    @staticmethod
    def _wishlist_model(model_key):
        key = (model_key or '').lower()
        if key not in WISHLIST_MODELS:
            raise ValidationError('Invalid wishlist item type.')
        app_label, model_name = WISHLIST_MODELS[key]
        try:
            return ContentType.objects.get(app_label=app_label, model=model_name)
        except ContentType.DoesNotExist as exc:
            raise ValidationError('Unsupported wishlist item.') from exc

    @staticmethod
    def add_wishlist(user, model_key, object_id):
        content_type = DashboardService._wishlist_model(model_key)
        item, created = WishlistItem.objects.get_or_create(
            user=user, content_type=content_type, object_id=object_id
        )
        return item, created

    @staticmethod
    def remove_wishlist(user, item_id):
        deleted, _ = WishlistItem.objects.filter(user=user, id=item_id).delete()
        return deleted

    @staticmethod
    def toggle_wishlist(user, model_key, object_id):
        """Add if absent, remove if present. Returns (action, item)."""
        content_type = DashboardService._wishlist_model(model_key)
        item = WishlistItem.objects.filter(
            user=user, content_type=content_type, object_id=object_id
        ).first()
        if item:
            item.delete()
            return 'removed', None
        item = WishlistItem.objects.create(
            user=user, content_type=content_type, object_id=object_id
        )
        return 'added', item

    @staticmethod
    def is_in_wishlist(user, model_key, object_id):
        try:
            content_type = DashboardService._wishlist_model(model_key)
        except ValidationError:
            return False
        return WishlistItem.objects.filter(
            user=user, content_type=content_type, object_id=object_id
        ).exists()

    # ------------------------------------------------------- preferences

    @staticmethod
    def update_profile(user, data):
        for field in ('first_name', 'last_name', 'phone_number',
                      'date_of_birth', 'address', 'city', 'country'):
            if field in data:
                value = data[field]
                setattr(user, field, value or None)
        user.save()
        return user

    @staticmethod
    def update_travel_preferences(user, data):
        profile = DashboardService.get_profile(user)
        profile.ensure_defaults()
        prefs = profile.travel_preferences
        allowed = UserProfile.DEFAULT_TRAVEL_PREFS.keys()
        for key in allowed:
            if key == 'hotel_amenities':
                prefs[key] = data.getlist(key) if hasattr(data, 'getlist') else data.get(key, [])
            elif key in data:
                prefs[key] = data[key]
        profile.travel_preferences = prefs
        profile.save(update_fields=['travel_preferences', 'updated_at'])
        return profile

    @staticmethod
    def update_notification_preferences(user, data):
        profile = DashboardService.get_profile(user)
        profile.ensure_defaults()
        prefs = profile.notification_preferences
        for key in UserProfile.DEFAULT_NOTIF_PREFS.keys():
            prefs[key] = bool(data.get(key))
        profile.notification_preferences = prefs
        profile.save(update_fields=['notification_preferences', 'updated_at'])
        return profile

    @staticmethod
    def change_password(user, old_password, new_password):
        if not user.check_password(old_password):
            raise ValidationError('Your current password is incorrect.')
        validate_password(new_password, user=user)
        user.set_password(new_password)
        user.save(update_fields=['password'])
        return user


def _bucket_dates(buckets):
    import datetime
    return [datetime.date(y, m, 1) for y, m in buckets.keys()]


# ===========================================================================
# ADMIN ANALYTICS
# Aggregations powering the staff-only executive dashboard. Views stay thin
# and every entry point is gated behind `is_staff` (covers superusers).
# ===========================================================================

RANGE_PRESETS = ('today', 'week', 'month', '6m', 'year', 'custom')


def _delta_class(delta, good_up):
    if delta is None:
        return 'is-neutral'
    is_good = delta >= 0 if good_up else delta < 0
    return 'is-up' if is_good else 'is-down'


def _delta_icon(delta):
    if delta is None:
        return 'bi-dash-lg'
    return 'bi-arrow-up-right' if delta >= 0 else 'bi-arrow-down-right'

ACTIVE_BOOKING_STATUSES = (
    Booking.Status.PENDING,
    Booking.Status.CONFIRMED,
    Booking.Status.COMPLETED,
)

CATEGORY_META = {
    Booking.BookingType.FLIGHT: ('Flights', 'bi-airplane-fill', '#06B6D4'),
    Booking.BookingType.HOTEL: ('Hotels', 'bi-building-fill', '#2563EB'),
    Booking.BookingType.PACKAGE: ('Packages', 'bi-box-seam-fill', '#8B5CF6'),
}


class AdminAnalyticsService:
    """Read-only analytical queries across the whole platform."""

    # ------------------------------------------------------------- filters

    @staticmethod
    def parse_range(query_dict):
        """Resolve the sticky-filter preset into an aware (start, end) window."""
        preset = (query_dict or {}).get('range', 'year')
        if preset not in RANGE_PRESETS:
            preset = 'year'
        now = timezone.now()
        today = now.date()

        def aware(d):
            return timezone.make_aware(datetime.combine(d, time.min))

        if preset == 'today':
            start, end, label = aware(today), now, 'Today'
        elif preset == 'week':
            start, end, label = (
                aware(today - timedelta(days=today.weekday())), now, 'This Week'
            )
        elif preset == 'month':
            start, end, label = aware(today.replace(day=1)), now, 'This Month'
        elif preset == '6m':
            start, end, label = now - timedelta(days=182), now, 'Last 6 Months'
        elif preset == 'custom':
            try:
                from_d = datetime.strptime(query_dict.get('from', ''), '%Y-%m-%d').date()
                to_d = datetime.strptime(query_dict.get('to', ''), '%Y-%m-%d').date()
                if from_d > to_d:
                    raise ValueError
                start = aware(from_d)
                end = aware(to_d + timedelta(days=1))
                label = f'{from_d:%b %d, %Y} — {to_d:%b %d, %Y}'
            except (ValueError, TypeError):
                start, end, label = aware(today.replace(month=1, day=1)), now, 'This Year'
                preset = 'year'
        else:  # year
            start, end, label = aware(today.replace(month=1, day=1)), now, 'This Year'

        return {
            'preset': preset,
            'label': label,
            'start': start,
            'end': end,
            'from': start.date().isoformat(),
            'to': (end - timedelta(microseconds=1)).date().isoformat(),
        }

    @staticmethod
    def _previous_window(window):
        span = window['end'] - window['start']
        prev_end = window['start']
        return {
            'preset': window['preset'],
            'label': window['label'],
            'start': prev_end - span,
            'end': prev_end,
            'from': (prev_end - span).date().isoformat(),
            'to': (prev_end - timedelta(microseconds=1)).date().isoformat(),
        }

    # ---------------------------------------------------------------- KPIs

    @staticmethod
    def _kpi_row(window):
        start, end = window['start'], window['end']

        bookings = Booking.objects.filter(created_at__range=(start, end))
        payments = Payment.objects.filter(created_at__range=(start, end))
        success = Payment.objects.filter(
            status=Payment.Status.SUCCESS,
            paid_at__range=(start, end),
        )
        revenue = success.aggregate(t=Sum('amount'))['t'] or Decimal('0')

        active_ids = set(bookings.values_list('user_id', flat=True))
        active_ids |= set(payments.values_list('user_id', flat=True))
        active_ids |= set(
            Review.objects.filter(created_at__range=(start, end))
            .values_list('user_id', flat=True)
        )

        return {
            'total_users': User.objects.count(),
            'new_users': User.objects.filter(
                date_joined__range=(start, end)
            ).count(),
            'active_users': len(active_ids),
            'total_bookings': bookings.count(),
            'total_revenue': float(revenue),
            'flights_booked': bookings.filter(
                booking_type=Booking.BookingType.FLIGHT
            ).count(),
            'hotels_booked': bookings.filter(
                booking_type=Booking.BookingType.HOTEL
            ).count(),
            'packages_booked': bookings.filter(
                booking_type=Booking.BookingType.PACKAGE
            ).count(),
            'pending_payments': payments.filter(
                status=Payment.Status.PENDING
            ).count(),
            'successful_payments': payments.filter(
                status=Payment.Status.SUCCESS
            ).count(),
            'cancelled_bookings': bookings.filter(
                status=Booking.Status.CANCELLED
            ).count(),
        }

    @staticmethod
    def _delta(cur, prev):
        if prev == 0:
            return 100.0 if cur > 0 else None
        return round((cur - prev) / prev * 100, 1)

    @staticmethod
    def kpi_cards(window):
        cur = AdminAnalyticsService._kpi_row(window)
        prev = AdminAnalyticsService._kpi_row(
            AdminAnalyticsService._previous_window(window)
        )

        specs = [
            ('total_users', 'Total Users', 'bi-people-fill', 'sky', True),
            ('active_users', 'Active Users', 'bi-person-check-fill', 'emerald', True),
            ('total_bookings', 'Total Bookings', 'bi-journal-bookmark-fill', 'violet', True),
            ('total_revenue', 'Total Revenue', 'bi-cash-stack', 'blue', True),
            ('flights_booked', 'Flights Booked', 'bi-airplane-fill', 'cyan', True),
            ('hotels_booked', 'Hotels Booked', 'bi-building-fill', 'indigo', True),
            ('packages_booked', 'Packages Booked', 'bi-box-seam-fill', 'purple', True),
            ('pending_payments', 'Pending Payments', 'bi-hourglass-split', 'amber', False),
            ('successful_payments', 'Successful Payments', 'bi-check-circle-fill', 'green', True),
            ('cancelled_bookings', 'Cancelled Bookings', 'bi-x-circle-fill', 'rose', False),
        ]

        cards = []
        for key, label, icon, variant, good_up in specs:
            value = cur[key]
            delta = AdminAnalyticsService._delta(cur[key], prev[key])
            is_money = key == 'total_revenue'
            cards.append({
                'key': key,
                'label': label,
                'icon': icon,
                'variant': variant,
                'good_up': good_up,
                'prefix': '₹' if is_money else '',
                'suffix': '',
                'decimals': 0,
                'value': float(value),
                'value_display': (
                    f'₹{float(value):,.0f}' if is_money
                    else f'{value:,}'
                ),
            'delta': delta,
            'delta_display': (
                '—' if delta is None else f'{"+" if delta >= 0 else ""}{delta}%'
            ),
            'delta_class': _delta_class(delta, good_up),
            'delta_icon': _delta_icon(delta),
        })
        return cards

    # -------------------------------------------------------------- charts

    @staticmethod
    def _month_keys(start, end):
        keys = []
        cur = start.date().replace(day=1)
        last = end.date().replace(day=1)
        while cur <= last:
            keys.append((cur.year, cur.month))
            y = cur.year + (1 if cur.month == 12 else 0)
            m = 1 if cur.month == 12 else cur.month + 1
            cur = cur.replace(year=y, month=m)
        return keys

    @staticmethod
    def _month_series(start, end):
        keys = AdminAnalyticsService._month_keys(start, end)
        labels = [datetime(y, m, 1).strftime('%b %y') for y, m in keys]
        return keys, labels

    @staticmethod
    def _monthly_totals(start, end, qs, date_field):
        keys, labels = AdminAnalyticsService._month_series(start, end)
        bucket = OrderedDict((k, Decimal('0')) for k in keys)
        for obj in qs:
            value = getattr(obj, date_field)
            if value is None:
                continue
            key = (value.year, value.month)
            if key in bucket:
                bucket[key] += obj.amount if date_field == 'paid_at' else 1
        return labels, [float(bucket[k]) for k in bucket]

    @staticmethod
    def revenue_by_month(start, end):
        qs = Payment.objects.filter(
            status=Payment.Status.SUCCESS,
            paid_at__range=(start, end),
        ).only('amount', 'paid_at')
        return AdminAnalyticsService._monthly_totals(start, end, qs, 'paid_at')

    @staticmethod
    def bookings_by_month(start, end):
        qs = Booking.objects.filter(created_at__range=(start, end)).only('created_at')
        return AdminAnalyticsService._monthly_totals(start, end, qs, 'created_at')

    @staticmethod
    def user_growth(start, end):
        qs = User.objects.filter(date_joined__range=(start, end)).only('date_joined')
        return AdminAnalyticsService._monthly_totals(start, end, qs, 'date_joined')

    @staticmethod
    def bookings_by_category(start, end):
        rows = (
            Booking.objects.filter(created_at__range=(start, end))
            .values('booking_type')
            .annotate(total=Count('id'))
        )
        result = {'labels': [], 'values': [], 'colors': []}
        for row in rows:
            if row['booking_type'] not in CATEGORY_META:
                continue
            name, _, color = CATEGORY_META[row['booking_type']]
            result['labels'].append(name)
            result['values'].append(row['total'])
            result['colors'].append(color)
        return result

    @staticmethod
    def popular_destinations(start, end, limit=7):
        qs = (
            Booking.objects.filter(
                created_at__range=(start, end),
                status__in=ACTIVE_BOOKING_STATUSES,
            )
            .select_related('flight', 'hotel', 'package')
        )
        counter = OrderedDict()
        for b in qs:
            dest = b.destination_display
            if not dest or dest == 'N/A':
                continue
            counter[dest] = counter.get(dest, 0) + 1
        top = sorted(counter.items(), key=lambda kv: kv[1], reverse=True)[:limit]
        return {'labels': [k for k, _ in top], 'values': [v for _, v in top]}

    @staticmethod
    def _top_items(start, end, booking_type, field, limit=6):
        rows = (
            Booking.objects.filter(
                created_at__range=(start, end),
                booking_type=booking_type,
                status__in=ACTIVE_BOOKING_STATUSES,
                **{f'{field}__isnull': False},
            )
            .values(field)
            .annotate(total=Count('id'))
            .order_by('-total')[:limit]
        )
        return {
            'labels': [r[field] for r in rows],
            'values': [r['total'] for r in rows],
        }

    @staticmethod
    def top_hotels(start, end, limit=6):
        return AdminAnalyticsService._top_items(
            start, end, Booking.BookingType.HOTEL, 'hotel__name', limit
        )

    @staticmethod
    def top_packages(start, end, limit=6):
        return AdminAnalyticsService._top_items(
            start, end, Booking.BookingType.PACKAGE, 'package__title', limit
        )

    @staticmethod
    def payment_health(start, end):
        rows = (
            Payment.objects.filter(created_at__range=(start, end))
            .values('status')
            .annotate(total=Count('id'))
        )
        counts = {r['status']: r['total'] for r in rows}
        return {
            'labels': ['Successful', 'Failed'],
            'values': [
                counts.get(Payment.Status.SUCCESS, 0),
                counts.get(Payment.Status.FAILED, 0),
            ],
            'colors': ['#22C55E', '#EF4444'],
        }

    @staticmethod
    def charts(window):
        start, end = window['start'], window['end']
        return {
            'revenue_by_month': dict(
                zip(
                    ('labels', 'values'),
                    AdminAnalyticsService.revenue_by_month(start, end),
                )
            ),
            'bookings_by_month': dict(
                zip(
                    ('labels', 'values'),
                    AdminAnalyticsService.bookings_by_month(start, end),
                )
            ),
            'user_growth': dict(
                zip(
                    ('labels', 'values'),
                    AdminAnalyticsService.user_growth(start, end),
                )
            ),
            'bookings_by_category': AdminAnalyticsService.bookings_by_category(start, end),
            'popular_destinations': AdminAnalyticsService.popular_destinations(start, end),
            'top_hotels': AdminAnalyticsService.top_hotels(start, end),
            'top_packages': AdminAnalyticsService.top_packages(start, end),
            'payment_health': AdminAnalyticsService.payment_health(start, end),
        }

    # ------------------------------------------------------------ activity

    @staticmethod
    def recent_users(limit=5):
        users = User.objects.order_by('-date_joined')[:limit]
        return [{
            'username': u.username,
            'name': u.get_full_name() or u.username,
            'email': u.email,
            'role': u.get_role_display() if u.role else 'Customer',
            'is_staff': u.is_staff,
            'date': u.date_joined,
            'initial': (u.username or '?')[:1].upper(),
        } for u in users]

    @staticmethod
    def recent_bookings(limit=5):
        qs = (
            Booking.objects.select_related('user', 'flight', 'hotel', 'package')
            .order_by('-created_at')[:limit]
        )
        return [{
            'booking_id': b.booking_id,
            'user': b.user.username,
            'type': b.booking_type,
            'title': b.title_display,
            'destination': b.destination_display,
            'amount': float(b.grand_total),
            'status': b.status,
            'color': b.status_color,
            'date': b.created_at,
            'detail_url': f'/bookings/{b.id}/',
        } for b in qs]

    @staticmethod
    def recent_payments(limit=5):
        qs = Payment.objects.select_related('user', 'booking').order_by('-created_at')[:limit]
        return [{
            'txn': p.transaction_id,
            'user': p.user.username,
            'booking_id': p.booking.booking_id,
            'amount': float(p.amount),
            'gateway': p.get_gateway_display(),
            'status': p.status,
            'color': p.status_color,
            'date': p.created_at,
        } for p in qs]

    @staticmethod
    def recent_reviews(limit=5):
        qs = Review.objects.select_related('user', 'hotel', 'package').order_by('-created_at')[:limit]
        return [{
            'user': r.user.username,
            'rating': r.rating,
            'title': r.title,
            'item': r.item_name,
            'approved': r.is_approved,
            'date': r.created_at,
            'detail_url': f'/reviews/{r.id}/',
        } for r in qs]

    @staticmethod
    def activity():
        return {
            'users': AdminAnalyticsService.recent_users(),
            'bookings': AdminAnalyticsService.recent_bookings(),
            'payments': AdminAnalyticsService.recent_payments(),
            'reviews': AdminAnalyticsService.recent_reviews(),
        }

    # -------------------------------------------------------------- search

    @staticmethod
    def search(query):
        q = (query or '').strip()
        if len(q) < 2:
            return {'query': q, 'users': [], 'bookings': [], 'payments': [], 'destinations': []}

        users = list(
            User.objects.filter(
                Q(username__icontains=q)
                | Q(email__icontains=q)
                | Q(first_name__icontains=q)
                | Q(last_name__icontains=q)
            ).order_by('-date_joined')[:4]
        )
        bookings = list(
            Booking.objects.select_related('user', 'hotel', 'package', 'flight')
            .filter(
                Q(booking_id__icontains=q)
                | Q(contact_name__icontains=q)
                | Q(contact_email__icontains=q)
            ).order_by('-created_at')[:4]
        )
        payments = list(
            Payment.objects.select_related('user', 'booking')
            .filter(
                Q(transaction_id__icontains=q)
                | Q(order_id__icontains=q)
                | Q(booking__booking_id__icontains=q)
            ).order_by('-created_at')[:4]
        )

        destinations = []
        for hotel in Hotel.objects.filter(
            Q(city__icontains=q) | Q(country__icontains=q) | Q(name__icontains=q)
        )[:3]:
            destinations.append({
                'label': f'{hotel.name}',
                'meta': f'{hotel.city}, {hotel.country} · Hotel',
                'url': f'/hotels/{hotel.slug}/',
            })
        for pkg in TravelPackage.objects.filter(
            Q(destination__icontains=q) | Q(country__icontains=q) | Q(title__icontains=q)
        )[:3]:
            destinations.append({
                'label': f'{pkg.destination}',
                'meta': f'{pkg.country} · Package',
                'url': f'/packages/{pkg.slug}/',
            })
        for fl in Flight.objects.filter(
            Q(departure_city__icontains=q) | Q(destination_city__icontains=q)
        )[:3]:
            destinations.append({
                'label': f'{fl.departure_city} → {fl.destination_city}',
                'meta': f'{fl.airline_name} {fl.flight_number} · Flight',
                'url': f'/flights/{fl.id}/',
            })

        return {
            'query': q,
            'users': [{
                'username': u.username,
                'name': u.get_full_name() or u.username,
                'email': u.email,
                'role': u.get_role_display() if u.role else 'Customer',
                'initial': (u.username or '?')[:1].upper(),
            } for u in users],
            'bookings': [{
                'booking_id': b.booking_id,
                'user': b.user.username,
                'type': b.booking_type,
                'destination': b.destination_display,
                'amount': float(b.grand_total),
                'status': b.status,
                'color': b.status_color,
                'date': b.created_at.isoformat(),
            } for b in bookings],
            'payments': [{
                'txn': p.transaction_id,
                'user': p.user.username,
                'booking_id': p.booking.booking_id,
                'amount': float(p.amount),
                'status': p.status,
                'color': p.status_color,
                'date': p.created_at.isoformat(),
            } for p in payments],
            'destinations': destinations,
        }

    # -------------------------------------------------------------- reports

    @staticmethod
    def report_rows(report, window):
        start, end = window['start'], window['end']

        if report == 'revenue':
            qs = (
                Payment.objects.select_related('user', 'booking')
                .filter(paid_at__range=(start, end))
                .order_by('-paid_at')
            )
            headers = ['Transaction ID', 'User', 'Booking ID', 'Amount', 'Currency',
                       'Gateway', 'Method', 'Status', 'Paid At']
            rows = [[
                p.transaction_id, p.user.username, p.booking.booking_id,
                float(p.amount), p.currency, p.get_gateway_display(),
                p.payment_method, p.status,
                p.paid_at.strftime('%d %b %Y %H:%M') if p.paid_at else '-',
            ] for p in qs]

        elif report == 'bookings':
            qs = (
                Booking.objects.select_related('user', 'flight', 'hotel', 'package')
                .filter(created_at__range=(start, end)).order_by('-created_at')
            )
            headers = ['Booking ID', 'User', 'Type', 'Destination', 'Travelers',
                       'Status', 'Payment Status', 'Total (INR)', 'Created At']
            rows = [[
                b.booking_id, b.user.username, b.booking_type,
                b.destination_display, b.total_travelers, b.status,
                b.payment_status, float(b.grand_total),
                b.created_at.strftime('%d %b %Y %H:%M'),
            ] for b in qs]

        elif report == 'payments':
            qs = (
                Payment.objects.select_related('user', 'booking')
                .filter(created_at__range=(start, end)).order_by('-created_at')
            )
            headers = ['Transaction ID', 'User', 'Booking ID', 'Amount', 'Currency',
                       'Gateway', 'Method', 'Status', 'Created At']
            rows = [[
                p.transaction_id, p.user.username, p.booking.booking_id,
                float(p.amount), p.currency, p.get_gateway_display(),
                p.payment_method, p.status,
                p.created_at.strftime('%d %b %Y %H:%M'),
            ] for p in qs]

        elif report == 'users':
            qs = (
                User.objects.filter(date_joined__range=(start, end))
                .annotate(
                    booking_count=Count('bookings', distinct=True),
                    spent=Sum(
                        'payments__amount',
                        filter=Q(payments__status=Payment.Status.SUCCESS),
                        distinct=True,
                    ),
                )
                .order_by('-date_joined')
            )
            headers = ['Username', 'Name', 'Email', 'Role', 'Verified',
                       'Bookings', 'Spent (INR)', 'Date Joined', 'Last Login']
            rows = [[
                u.username, u.get_full_name(), u.email,
                u.get_role_display() if u.role else 'Customer',
                'Yes' if u.is_email_verified else 'No',
                u.booking_count,
                float(u.spent or 0),
                u.date_joined.strftime('%d %b %Y %H:%M'),
                u.last_login.strftime('%d %b %Y %H:%M') if u.last_login else '-',
            ] for u in qs]

        else:
            raise ValueError('Unknown report')

        return headers, rows

    @staticmethod
    def export(report, fmt, window):
        headers, rows = AdminAnalyticsService.report_rows(report, window)
        slug = f'{report}_report_{window["start"].date():%Y-%m-%d}_{window["end"].date():%Y-%m-%d}'
        if fmt == 'csv':
            return AdminAnalyticsService._export_csv(headers, rows, slug)
        if fmt == 'xlsx':
            return AdminAnalyticsService._export_xlsx(headers, rows, slug)
        return AdminAnalyticsService._export_pdf(report, headers, rows, window, slug)

    @staticmethod
    def _export_csv(headers, rows, slug):
        import csv
        import io
        buf = io.StringIO()
        buf.write('\ufeff')  # BOM so Excel opens UTF-8 correctly
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse(buf.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{slug}.csv"'
        return response

    @staticmethod
    def _export_xlsx(headers, rows, slug):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = 'Report'
        ws.append(headers)
        header_fill = PatternFill('solid', fgColor='2563EB')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(vertical='center')
        for row in rows:
            ws.append(row)
        for i, _ in enumerate(headers, start=1):
            widths = {
                'A': 20, 'B': 18, 'C': 18, 'D': 14, 'E': 12,
                'F': 14, 'G': 12, 'H': 16, 'I': 20,
            }
            ws.column_dimensions[get_column_letter(i)].width = widths.get(get_column_letter(i), 16)
        ws.freeze_panes = 'A2'

        from io import BytesIO
        buf = BytesIO()
        wb.save(buf)
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{slug}.xlsx"'
        return response

    @staticmethod
    def _export_pdf(report, headers, rows, window, slug):
        from io import BytesIO

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=landscape(A4),
            rightMargin=0.5 * inch, leftMargin=0.5 * inch,
            topMargin=0.55 * inch, bottomMargin=0.55 * inch,
            title=f'{report.title()} Report',
            author='TravelBooking',
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'AnalyticsTitle', parent=styles['Title'],
            fontSize=16, leading=20, textColor=colors.HexColor('#0F172A'),
            spaceAfter=2,
        )
        meta_style = ParagraphStyle(
            'AnalyticsMeta', parent=styles['Normal'],
            fontSize=8.5, leading=12, textColor=colors.HexColor('#64748B'),
            spaceAfter=12,
        )
        header_style = ParagraphStyle(
            'AnalyticsHead', parent=styles['Normal'],
            fontSize=8, leading=10, textColor=colors.white,
            fontName='Helvetica-Bold',
        )
        cell_style = ParagraphStyle(
            'AnalyticsCell', parent=styles['Normal'],
            fontSize=8, leading=10, textColor=colors.HexColor('#334155'),
        )

        story = [
            Paragraph(f'{report.title()} Report', title_style),
            Paragraph(
                f'Period: {window["label"]} · '
                f'Generated {timezone.now():%d %b %Y, %H:%M} · TravelBooking Analytics',
                meta_style,
            ),
            Spacer(1, 4),
        ]

        data = [[Paragraph(escape(h), header_style) for h in headers]]
        data += [[Paragraph(escape(str(c)), cell_style) for c in row] for row in rows]
        table = Table(data, repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1),
             [colors.white, colors.HexColor('#F1F5F9')]),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#E2E8F0')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(table)
        doc.build(story)

        response = HttpResponse(buf.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{slug}.pdf"'
        return response
